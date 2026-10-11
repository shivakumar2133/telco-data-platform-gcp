"""SLA targets: a business-owned Excel workbook (deliberately a bit messy)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from data_generator.defects import DefectLog

# (response_hours, resolution_hours) per tier and priority
SLA_MATRIX: dict[str, dict[str, tuple[float, float]]] = {
    "Premium": {"P1": (0.25, 4), "P2": (0.5, 8), "P3": (2, 24), "P4": (4, 72)},
    "Enhanced": {"P1": (0.5, 6), "P2": (1, 12), "P3": (4, 48), "P4": (8, 96)},
    "Standard": {"P1": (1, 12), "P2": (2, 24), "P3": (8, 72), "P4": (24, 120)},
}
AVAILABILITY = {"Premium": 99.99, "Enhanced": 99.95, "Standard": 99.9}
HEADERS = [
    "SLA Tier",
    "Priority",
    "Response (hrs)",
    "Resolution (hrs)",
    "Availability Target %",
    "Effective From",
]


def write_sla_workbook(path: Path, log: DefectLog) -> int:
    sys = "SLA workbook (XLSX)"
    rows: list[list] = []
    for tier, prios in SLA_MATRIX.items():
        for prio, (resp, res) in prios.items():
            rows.append([tier, prio, resp, res, AVAILABILITY[tier], date(2026, 1, 1)])

    # Defects a human-maintained spreadsheet typically has
    for r in rows:
        if r[0] == "Enhanced" and r[1] == "P3":
            r[0] = "Enhanced "  # trailing space
        if r[0] == "Standard" and r[1] == "P2":
            r[3] = "24 hrs"  # text in a numeric column
        if r[0] == "Standard" and r[1] == "P4":
            r[4] = 0.999  # fraction instead of percentage
    rows.insert(1, list(rows[0]))  # Premium P1 entered twice
    log.add("sla_trailing_whitespace", sys, "Tier name has trailing space", ["Enhanced/P3"])
    log.add("sla_text_in_numeric", sys, "Resolution hours stored as text '24 hrs'", ["Standard/P2"])
    log.add(
        "sla_inconsistent_percentage", sys, "Availability 0.999 instead of 99.9", ["Standard/P4"]
    )
    log.add("sla_duplicate_row", sys, "Same tier/priority row appears twice", ["Premium/P1"])

    wb = Workbook()
    ws = wb.active
    ws.title = "SLA Targets"
    ws.cell(row=1, column=1, value="Telco SLA Matrix (synthetic) - owned by Service Management")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    # Row 2 left blank on purpose: the header is on row 3, not row 1
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F46E5")
    for i, r in enumerate(rows, start=4):
        for col, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=col, value=v)
            if col == 6:
                c.number_format = "yyyy-mm-dd"
    for col, width in zip("ABCDEF", [14, 10, 16, 18, 22, 16], strict=True):
        ws.column_dimensions[col].width = width

    notes = wb.create_sheet("Notes")
    for i, line in enumerate(
        [
            "Response = time to acknowledge; Resolution = time to restore service.",
            "Hours are clock hours (24x7) for all tiers.",
            "Changes must be approved by the Service Management lead.",
        ],
        start=1,
    ):
        notes.cell(row=i, column=1, value=line)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows)
