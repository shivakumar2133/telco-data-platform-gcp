import io
import json

import pytest
from openpyxl import Workbook

from ingestion.common.config import Settings
from ingestion.common.context import RunContext
from ingestion.common.lake import LocalLake
from ingestion.common.metadata import LocalBronzeLoader, LocalMetaStore
from ingestion.extractors.file_processor import (
    CDR_COLUMNS,
    FileRejected,
    parse_alarms,
    parse_cdr,
    parse_sla,
    process_landing,
)

GOOD_CDR = (
    ",".join(CDR_COLUMNS) + "\n"
    "CDR1,447000000001,447000000002,VOICE,2026-09-01 10:00:00,-5,0.0,CELL1,COMPLETED,0.05\n"
)


def alarms_doc(count_in_header: int, n: int = 2) -> bytes:
    alarms = [
        {
            "alarm_id": f"A{i}",
            "device": {"device_id": "D1", "site_id": "S1"},
            "severity": "MAJOR",
            "attributes": {"probable_cause": "x"},
        }
        for i in range(n)
    ]
    return json.dumps(
        {
            "export_id": "E1",
            "exported_at": "2026-09-02T01:00:00Z",
            "record_count": count_in_header,
            "alarms": alarms,
        }
    ).encode()


def sla_xlsx(with_header: bool = True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SLA Targets"
    ws["A1"] = "Title row"
    if with_header:
        for i, h in enumerate(["SLA Tier", "Priority", "Response (hrs)"], start=1):
            ws.cell(row=3, column=i, value=h)
        ws.append(["Premium", "P1", 0.25])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_cdr_keeps_bad_rows_as_text():
    df = parse_cdr(GOOD_CDR.encode())
    assert df.loc[0, "duration_sec"] == "-5"  # row-level issues are Silver's job


def test_cdr_wrong_header_rejected():
    with pytest.raises(FileRejected, match="header"):
        parse_cdr(b"id,when\n1,2\n")


def test_alarms_flattened_with_header_fields():
    df = parse_alarms(alarms_doc(2))
    assert {"device__device_id", "attributes__probable_cause", "export_id"} <= set(df.columns)


def test_alarms_control_total_mismatch_rejected():
    with pytest.raises(FileRejected, match="control total"):
        parse_alarms(alarms_doc(3, n=2))


def test_sla_finds_header_on_row_3():
    df = parse_sla(sla_xlsx())
    assert list(df.columns) == ["sla_tier", "priority", "response_hrs"]
    assert df.loc[0, "sla_tier"] == "Premium"


def test_sla_without_header_rejected():
    with pytest.raises(FileRejected):
        parse_sla(sla_xlsx(with_header=False))


@pytest.fixture
def ctx(tmp_path):
    s = Settings(project_id="p", region="us-central1", bq_location="us-central1", env="test")
    return RunContext(
        settings=s,
        lake=LocalLake(tmp_path / "lake"),
        meta=LocalMetaStore(tmp_path / "meta"),
        loader=LocalBronzeLoader(),
        ingest_sa="",
    )


def test_process_landing_end_to_end(ctx):
    ctx.lake.write_bytes("landing/cdr/cdr_20260901.csv", GOOD_CDR.encode())
    ctx.lake.write_bytes("landing/alarms/alarms_20260901.json", alarms_doc(9))  # bad total
    results = {r.object_name: r for r in process_landing(ctx)}

    assert results["cdr/cdr_20260901.csv"].status == "SUCCESS"
    assert results["cdr/cdr_20260901.csv"].loaded_row_count == 1
    assert results["alarms/alarms_20260901.json"].status == "REJECTED"
    assert ctx.lake.exists("archive/cdr/cdr_20260901.csv")
    assert ctx.lake.exists("quarantine/alarms/alarms_20260901.json")
    assert ctx.lake.list("landing/") == []

    # Same file delivered again -> skipped, not loaded twice
    ctx.lake.write_bytes("landing/cdr/cdr_20260901.csv", GOOD_CDR.encode())
    again = process_landing(ctx)
    assert [r.status for r in again] == ["SKIPPED"]
